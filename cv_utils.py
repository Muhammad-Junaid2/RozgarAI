"""
cv_utils.py
Utilities for extracting plain text from an uploaded CV file (PDF, DOCX, or TXT).
"""
from __future__ import annotations
import io


class CVExtractionError(Exception):
    """Raised when text cannot be extracted from the uploaded file."""


def extract_text_from_pdf(file_bytes: bytes) -> str:
    import pdfplumber

    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    text = "\n".join(text_parts).strip()
    if not text:
        raise CVExtractionError(
            "No selectable text found in this PDF. It may be a scanned image — "
            "try exporting your CV as a text-based PDF or DOCX."
        )
    return text


def extract_text_from_docx(file_bytes: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    text = "\n".join(p for p in parts if p.strip()).strip()
    if not text:
        raise CVExtractionError("The DOCX file appears to be empty.")
    return text


def extract_text_from_txt(file_bytes: bytes) -> str:
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="ignore")
    text = text.strip()
    if not text:
        raise CVExtractionError("The text file appears to be empty.")
    return text


def extract_cv_text(filename: str, file_bytes: bytes) -> str:
    """Dispatch to the right extractor based on file extension."""
    name = filename.lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    if name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    if name.endswith(".txt"):
        return extract_text_from_txt(file_bytes)
    raise CVExtractionError(
        f"Unsupported file type for '{filename}'. Please upload a PDF, DOCX, or TXT file."
    )


def truncate_text(text: str, max_chars: int = 12000) -> str:
    """Keep prompts within a sane size for the LLM."""
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n[...CV truncated for length...]"
