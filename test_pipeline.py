"""
Quick offline test harness:
1. Runs the full career_manager_pipeline in MOCK MODE (no API key needed) and
   validates every expected key exists in the output.
2. Builds a real PDF and a real DOCX on the fly and checks cv_utils extracts
   text from them correctly.
"""
import os
os.environ["MOCK_MODE"] = "1"
os.environ.pop("GEMINI_API_KEY", None)

from agents import career_manager_pipeline  # noqa: E402
from cv_utils import extract_cv_text, extract_text_from_txt, CVExtractionError  # noqa: E402

PASS = "PASS"
FAIL = "FAIL"


def check(label, condition):
    print(f"[{PASS if condition else FAIL}] {label}")
    return condition


def test_pipeline():
    with open("sample_data/sample_cv.txt") as f:
        cv_text = f.read()
    with open("sample_data/sample_jd.txt") as f:
        jd_text = f.read()

    log = []
    results = career_manager_pipeline(cv_text, jd_text, progress_cb=lambda s: log.append(s))

    all_ok = True
    all_ok &= check("7 progress steps logged", len(log) == 7)

    expected_top_keys = {
        "cv_analysis", "skills", "job_match", "interview", "roadmap",
        "cover_letter", "linkedin",
    }
    all_ok &= check("all top-level keys present", expected_top_keys.issubset(results.keys()))

    ca = results["cv_analysis"]
    all_ok &= check(
        "cv_analysis has ats_score (int 0-100)",
        isinstance(ca.get("ats_score"), int) and 0 <= ca["ats_score"] <= 100,
    )
    all_ok &= check("cv_analysis has improvement_suggestions list", isinstance(ca.get("improvement_suggestions"), list))

    sk = results["skills"]
    all_ok &= check("skills has matching_skills list", isinstance(sk.get("matching_skills"), list))
    all_ok &= check("skills has missing_skills list", isinstance(sk.get("missing_skills"), list))

    jm = results["job_match"]
    all_ok &= check("job_match has verdict string", isinstance(jm.get("verdict"), str))

    iv = results["interview"]
    all_ok &= check(
        "interview has 6 questions with required fields",
        isinstance(iv.get("questions"), list)
        and len(iv["questions"]) == 6
        and all({"question", "type", "personalized_answer_outline"} <= set(q.keys()) for q in iv["questions"]),
    )

    rm = results["roadmap"]
    all_ok &= check(
        "roadmap has 4 weeks",
        isinstance(rm.get("roadmap"), list) and len(rm["roadmap"]) == 4,
    )

    cl = results["cover_letter"]
    all_ok &= check("cover_letter has non-empty text", bool(cl.get("cover_letter", "").strip()))

    li = results["linkedin"]
    all_ok &= check("linkedin has headline_options list", isinstance(li.get("headline_options"), list))

    return all_ok


def test_cv_extraction():
    all_ok = True

    # TXT
    txt_bytes = b"John Doe\nPython Developer\nSkills: Python, SQL"
    text = extract_text_from_txt(txt_bytes)
    all_ok &= check("txt extraction works", "Python Developer" in text)

    # Empty txt should raise
    try:
        extract_text_from_txt(b"   ")
        all_ok &= check("empty txt raises CVExtractionError", False)
    except CVExtractionError:
        all_ok &= check("empty txt raises CVExtractionError", True)

    # Build and test a real PDF
    try:
        from reportlab.pdfgen import canvas
        pdf_path = "/tmp/test_cv.pdf"
        c = canvas.Canvas(pdf_path)
        c.drawString(100, 750, "Jane Smith - Backend Developer")
        c.drawString(100, 730, "Skills: Flask, PostgreSQL, Docker")
        c.save()
        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()
        pdf_text = extract_cv_text("cv.pdf", pdf_bytes)
        all_ok &= check("pdf extraction works", "Backend Developer" in pdf_text and "Flask" in pdf_text)
    except ImportError:
        print("[SKIP] reportlab not installed, skipping real PDF test")

    # Build and test a real DOCX
    import docx
    doc = docx.Document()
    doc.add_paragraph("Ali Raza - Data Analyst")
    doc.add_paragraph("Skills: Pandas, NumPy, SQL")
    docx_path = "/tmp/test_cv.docx"
    doc.save(docx_path)
    with open(docx_path, "rb") as f:
        docx_bytes = f.read()
    docx_text = extract_cv_text("cv.docx", docx_bytes)
    all_ok &= check("docx extraction works", "Data Analyst" in docx_text and "Pandas" in docx_text)

    # Unsupported type should raise
    try:
        extract_cv_text("cv.exe", b"whatever")
        all_ok &= check("unsupported file type raises", False)
    except CVExtractionError:
        all_ok &= check("unsupported file type raises", True)

    return all_ok


if __name__ == "__main__":
    print("=== Pipeline test ===")
    ok1 = test_pipeline()
    print("\n=== CV extraction test ===")
    ok2 = test_cv_extraction()

    print("\n=== SUMMARY ===")
    if ok1 and ok2:
        print("ALL TESTS PASSED")
    else:
        print("SOME TESTS FAILED")
        raise SystemExit(1)
